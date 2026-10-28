"""Download and reopen exported models outside the service artifact directories."""
import argparse
import hashlib
import io
import json
import os
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import httpx

ROOT=Path(__file__).resolve().parent.parent


def unpack(content,directory):
    directory.mkdir(parents=True,exist_ok=False)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for member in archive.infolist():
            if not (directory/member.filename).resolve().is_relative_to(directory.resolve()):
                raise ValueError("Archive path escapes verification directory")
        archive.extractall(directory)


def run(script,directory,timeout=90):
    completed=subprocess.run([sys.executable,"-X","utf8",script],cwd=directory,
                             capture_output=True,text=True,encoding="utf-8",timeout=timeout)
    if completed.returncode:
        raise RuntimeError(f"{script} failed: "+(completed.stdout+completed.stderr)[-3000:])


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--base-url",default="http://127.0.0.1:8000")
    parser.add_argument("--study",type=Path,default=ROOT/"data/cantera-study-results.json")
    args=parser.parse_args()
    study=json.loads(args.study.read_text(encoding="utf-8"))
    stamp=datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S-%f")
    destination=ROOT/"data/qa"/("cantera-replay-"+stamp)
    destination.mkdir(parents=True,exist_ok=False)
    headers={"Authorization":"Bearer "+os.environ["CHEMELLM_AUTH_TOKEN"]} if os.environ.get("CHEMELLM_AUTH_TOKEN") else {}
    native=[]
    with httpx.Client(base_url=args.base_url,headers=headers,timeout=120) as client:
        def download(artifact):
            response=client.get(artifact["download_url"]);response.raise_for_status()
            if hashlib.sha256(response.content).hexdigest()!=artifact["metadata"]["sha256"]:
                raise ValueError("Download integrity failed: "+artifact["name"])
            return response.content
        for index,case in enumerate(study["cases"]):
            result=case["result"]
            archive=next(a for a in result["artifacts"] if a["id"]==result["native_archive_id"])
            folder=destination/("native-"+str(index+1))
            unpack(download(archive),folder)
            run("simulation.py",folder)
            replay=json.loads((folder/"native-replay-results.json").read_text(encoding="utf-8"))
            if not replay["passed"]:raise ValueError("Native replay failed")
            native.append({"case":case["label"],"directory":str(folder),"checks":replay["checks"],"passed":replay["passed"]})
        bundle_folder=destination/"project"
        unpack(download(study["bundle"]),bundle_folder)
    manifest=json.loads((bundle_folder/"manifest.json").read_text(encoding="utf-8"))
    integrity=[]
    for artifact in manifest["artifacts"]:
        path=bundle_folder/"files"/artifact["id"]/artifact["name"]
        actual=hashlib.sha256(path.read_bytes()).hexdigest()
        integrity.append({"id":artifact["id"],"name":artifact["name"],"passed":actual==artifact["metadata"]["sha256"]})
    if not all(a["passed"] for a in integrity):raise ValueError("Bundle artifact hash mismatch")
    (bundle_folder/"artifact-integrity.json").write_text(json.dumps({"passed":True,"files":integrity},indent=2),encoding="utf-8")
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfReader
    import pypdfium2 as pdfium
    documents={}
    for artifact in study["report"]["artifacts"]:
        path=bundle_folder/"files"/artifact["id"]/artifact["name"]
        if artifact["kind"]=="docx":
            document=Document(path)
            assert "Cantera" in "\n".join(p.text for p in document.paragraphs)
            documents["docx_paragraphs"]=len(document.paragraphs)
        elif artifact["kind"]=="pptx":
            presentation=Presentation(path)
            assert len(presentation.slides)>0
            documents["pptx_slides"]=len(presentation.slides)
        elif artifact["kind"]=="pdf":
            reader=PdfReader(path)
            assert "Cantera" in "\n".join(p.extract_text() for p in reader.pages)
            documents["pdf_pages"]=len(reader.pages)
            document=pdfium.PdfDocument(path)
            for index in range(len(document)):
                page=document[index];bitmap=page.render(scale=1.2)
                bitmap.to_pil().save(destination/(f"report-page-{index+1}.png"))
                bitmap.close();page.close()
            document.close()
        elif artifact["kind"]=="md":
            assert "Cantera" in path.read_text(encoding="utf-8")
            documents["markdown_readable"]=True
    run("replay.py",bundle_folder/"reproduce")
    replay=json.loads((bundle_folder/"reproduce/replay-results.json").read_text(encoding="utf-8"))
    if len(replay["results"])!=len(study["cases"]) or not all(a["reproduced"] for a in replay["results"]):
        raise ValueError("Generic bundle replay count/result failed")
    report={"passed":True,"project_id":study["project"]["id"],"directory":str(destination),"native":native,
            "bundle_files_verified":len(integrity),"generic_calculations_reproduced":len(replay["results"]),"documents":documents,
            "scope":"Export integrity and same-engine native/source replay; not independent physical, LLM or Linux validation."}
    (destination/"verification.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    (ROOT/"data/qa/cantera-study-verification.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))


if __name__=="__main__":main()
