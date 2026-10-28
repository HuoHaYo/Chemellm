import io
import json
import zipfile
from pathlib import Path
import ezdxf
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from app.config import ROOT
from app.db import one, execute
from app.diagrams import DEMO

def test_competition_upload_and_evidence_gate(client,project_id):
    data=(ROOT/"26化工设计大赛题目.pdf").read_bytes()
    response=client.post(f"/api/projects/{project_id}/upload",files={"file":("26化工设计大赛题目.pdf",data,"application/pdf")})
    assert response.status_code==200,response.text
    record=response.json()
    assert len(record["metadata"]["pages"])==5
    assert record["metadata"]["competition_match"]
    acceptance=client.get(f"/api/projects/{project_id}/acceptance").json()
    assert acceptance["total"]==35 and acceptance["mandatory"]==31
    assert not acceptance["ready"] and acceptance["verified"]==0
    rejected=client.put(f"/api/projects/{project_id}/requirements/D01",json={"status":"verified","notes":"已审核"})
    assert rejected.status_code==422
    preview=client.get(record["preview_url"])
    assert preview.content.startswith(b"%PDF")
    assert client.get(record["download_url"]).content==data
    results=client.post(f"/api/projects/{project_id}/tools/project_search",json={"arguments":{"query":"gbcbig.shx"}}).json()
    assert results["results"][0]["page"]==5

def test_office_pdf_generation_and_preview(client,project_id):
    r=client.post(f"/api/projects/{project_id}/tools/export_report",json={"arguments":{"title":"设计校核记录","body":"苯乙烯设计依据\n质量闭合需独立检查。\n<script>alert(1)</script>","formats":["pdf","docx","pptx"]}})
    assert r.status_code==200,r.text
    for item in r.json()["artifacts"]:
        content=client.get(item["download_url"]).content
        if item["kind"]=="pdf":
            assert "设计" in "".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)
        if item["kind"]=="docx":
            assert Document(io.BytesIO(content)).paragraphs
        if item["kind"]=="pptx":
            assert len(Presentation(io.BytesIO(content)).slides)>0
        preview=client.get(item["preview_url"])
        assert preview.status_code==200
        if "text/html" in preview.headers["content-type"]:
            assert "<script>alert" not in preview.text
    assert len(client.get(f"/api/projects/{project_id}/artifacts").json())==3

def test_editable_drawing_exports_and_invalid_topology(client,project_id):
    r=client.post(f"/api/projects/{project_id}/tools/pfd",json={"arguments":DEMO})
    assert r.status_code==200,r.text
    kinds={a["kind"] for a in r.json()["artifacts"]}
    assert kinds=={"svg","pdf","json","dxf"}
    dxf=next(a for a in r.json()["artifacts"] if a["kind"]=="dxf")
    path=one("SELECT path FROM artifacts WHERE id=?",(dxf["id"],))["path"]
    doc=ezdxf.readfile(path)
    assert doc.dxfversion=="AC1018" and not doc.audit().has_errors
    rectangles=sum(n["kind"] not in {"pump","exchanger","instrument"} for n in DEMO["nodes"])
    assert len(doc.modelspace().query("LWPOLYLINE"))==len(DEMO["edges"])+rectangles
    invalid={**DEMO,"edges":[{"source":"no-such-unit","target":"R-101"}]}
    assert client.post(f"/api/projects/{project_id}/tools/pfd",json={"arguments":invalid}).status_code==422

def test_upload_rejections_and_auth(client,project_id,monkeypatch):
    assert client.post(f"/api/projects/{project_id}/upload",files={"file":("evil.svg",b"<svg><script/></svg>")}).status_code==422
    assert client.post(f"/api/projects/{project_id}/upload",files={"file":("bad.pdf",b"not pdf")}).status_code==422
    from app import config
    monkeypatch.setattr(config,"AUTH_TOKEN","test-access-token")
    assert client.get("/api/projects").status_code==401
    assert client.post("/api/login",json={"token":"bad"}).status_code==401
    login=client.post("/api/login",json={"token":"test-access-token"})
    assert login.status_code==200 and "httponly" in login.headers["set-cookie"].lower()
    assert client.get("/api/projects").status_code==200
    assert client.post("/api/projects",json={"name":"x"},headers={"Origin":"https://attacker.example"}).status_code==403

def test_no_model_does_not_fake_results(client,project_id,monkeypatch):
    monkeypatch.delenv("MODEL_BALANCED",raising=False)
    r=client.post(f"/api/projects/{project_id}/runs",json={"message":"模拟苯乙烯流程"})
    assert r.status_code==422
    assert "未配置" in r.json()["detail"]

def test_bundle_reproduces_recorded_calculation(client,project_id,tmp_path):
    import subprocess
    import sys
    assert client.post(f"/api/projects/{project_id}/tools/material_balance",json={"arguments":{"conversion":0.7}}).status_code==200
    assert client.post(f"/api/projects/{project_id}/tools/column_mesh",json={"arguments":{}}).status_code==200
    assert client.post(f"/api/projects/{project_id}/tools/aqueous_flash",json={"arguments":{}}).status_code==200
    assert client.post(f"/api/projects/{project_id}/tools/frontend_process",json={"arguments":{}}).status_code==200
    r=client.post(f"/api/projects/{project_id}/bundle")
    assert r.status_code==200,r.text
    data=client.get(r.json()["download_url"]).content
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "reproduce/replay.py" in archive.namelist()
        archive.extractall(tmp_path)
    result=subprocess.run([sys.executable,"-X","utf8","replay.py"],cwd=tmp_path/"reproduce",capture_output=True,text=True,encoding="utf-8",timeout=60)
    assert result.returncode==0,result.stdout+result.stderr
    output=json.loads((tmp_path/"reproduce"/"replay-results.json").read_text(encoding="utf-8"))
    assert len(output["results"])==4 and all(r["reproduced"] for r in output["results"])

def test_bundle_keeps_all_calculations_and_list_is_paginated(client,project_id):
    from app.db import uid,now
    for number in range(103):
        execute("INSERT INTO calculations VALUES(?,?,?,?,?,?)",(uid(),project_id,"convert_units",json.dumps({"value":number,"from_unit":"bar","to_unit":"Pa"}),json.dumps({"value":number*100000}),now()))
    first=client.get(f"/api/projects/{project_id}/calculations").json()
    rest=client.get(f"/api/projects/{project_id}/calculations?offset=100").json()
    assert len(first)==100 and len(rest)==3
    assert not {r["id"] for r in first}&{r["id"] for r in rest}
    r=client.post(f"/api/projects/{project_id}/bundle")
    with zipfile.ZipFile(io.BytesIO(client.get(r.json()["download_url"]).content)) as archive:
        manifest=json.loads(archive.read("manifest.json"))
    assert len(manifest["calculations"])==103
    assert client.get(f"/api/projects/{project_id}/calculations?limit=0").status_code==422
