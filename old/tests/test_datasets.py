import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from app.benchmarks import example_csv
from app.db import one, rows, execute, uid, now


def upload_csv(client, project_id, content=None):
    content = content if content is not None else example_csv("BoxBOD")[0]
    response = client.post(f"/api/projects/{project_id}/upload", files={"file":("experiment.csv",content)})
    assert response.status_code == 200, response.text
    return response.json()


def arguments(artifact):
    return {"artifact_id":artifact["id"],"x_column":"time_day","y_column":"BOD_mg_liter","x_unit":"day","y_unit":"mg/liter","models":["linear","first_order_uptake"]}


def test_analysis_source_comparison_and_exported_artifacts(client, project_id):
    artifact = upload_csv(client, project_id)
    inspected = client.get(f'/api/projects/{project_id}/datasets/{artifact["id"]}').json()
    assert inspected["rows"] == 6 and inspected["columns"] == ["time_day","BOD_mg_liter"]
    response = client.post(f"/api/projects/{project_id}/tools/analyze_dataset",json={"arguments":arguments(artifact)})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["source"]["sha256"] == artifact["metadata"]["sha256"]
    assert len(result["fits"]) == 2 and not result["failures"]
    assert {a["kind"] for a in result["artifacts"]} == {"svg","pdf","csv","json"}
    assert len(rows("SELECT id FROM calculations WHERE project_id=? AND tool='fit_data'",(project_id,))) == 2
    for generated in result["artifacts"]:
        assert client.get(generated["download_url"]).status_code == 200
        assert client.get(generated["preview_url"]).status_code == 200


def test_missing_values_and_formulas_not_imputed(client, project_id):
    artifact = upload_csv(client, project_id, b"time_day,BOD_mg_liter\n1,2\n2,\n3,4\n4,5\n")
    response = client.post(f"/api/projects/{project_id}/tools/analyze_dataset",json={"arguments":arguments(artifact)})
    assert response.status_code == 422 and "第3行" in response.text
    book = Workbook()
    book.active.append(["x","y"])
    for x, y in [(1,2),(2,"=A3*2"),(3,6),(4,8)]:
        book.active.append([x,y])
    book.create_sheet("measurements").append(["other_x","other_y"])
    stream = io.BytesIO(); book.save(stream)
    uploaded = client.post(f"/api/projects/{project_id}/upload",files={"file":("data.xlsx",stream.getvalue())})
    assert uploaded.status_code == 200, uploaded.text
    ident = uploaded.json()["id"]
    info = client.get(f"/api/projects/{project_id}/datasets/{ident}").json()
    assert info["sheets"] == ["Sheet","measurements"]
    response = client.post(f"/api/projects/{project_id}/tools/analyze_dataset",json={"arguments":{"artifact_id":ident,"x_column":"x","y_column":"y"}})
    assert response.status_code == 422 and "公式" in response.text


def test_excel_numeric_sheet_selection(client, project_id):
    book = Workbook(); book.active.append(["cover"]); book.active.append(["metadata"])
    sheet = book.create_sheet("data")
    sheet.append(["time","concentration"])
    for x, y in [(0,10),(1,8.2),(2,6.7),(3,5.5),(4,4.5)]:
        sheet.append([x,y])
    stream=io.BytesIO();book.save(stream)
    artifact=client.post(f"/api/projects/{project_id}/upload",files={"file":("study.xlsx",stream.getvalue())}).json()
    response=client.post(f"/api/projects/{project_id}/tools/analyze_dataset",json={"arguments":{"artifact_id":artifact["id"],"sheet":"data","x_column":"time","y_column":"concentration","x_unit":"s","y_unit":"mol/liter","models":["first_order_decay"]}})
    assert response.status_code == 200, response.text
    assert response.json()["source"]["sheet"] == "data"


def test_project_isolation_and_changed_source_rejection(client, project_id):
    artifact=upload_csv(client,project_id)
    other=uid();execute("INSERT INTO projects VALUES(?,?,?)",(other,"other",now()))
    assert client.get(f'/api/projects/{other}/datasets/{artifact["id"]}').status_code == 422
    row=one("SELECT path FROM artifacts WHERE id=?",(artifact["id"],))
    Path(row["path"]).write_text("x,y\n1,2\n",encoding="utf-8")
    response=client.get(f'/api/projects/{project_id}/datasets/{artifact["id"]}')
    assert response.status_code == 422 and "哈希" in response.text


def test_failed_candidate_is_visible_not_silently_removed(client, project_id):
    artifact=upload_csv(client,project_id,b"time_day,BOD_mg_liter\n0,0\n1,2\n2,3\n3,4\n4,5\n")
    args={**arguments(artifact),"models":["linear","power_law"]}
    response=client.post(f"/api/projects/{project_id}/tools/analyze_dataset",json={"arguments":args})
    assert response.status_code == 200, response.text
    assert response.json()["failures"][0]["model"] == "power_law"
    assert len(response.json()["fits"]) == 1


def test_experiment_bundle_replays_in_separate_process(client, project_id, tmp_path):
    artifact=upload_csv(client,project_id)
    assert client.post(f"/api/projects/{project_id}/tools/analyze_dataset",json={"arguments":arguments(artifact)}).status_code == 200
    response=client.post(f"/api/projects/{project_id}/bundle").json()
    with zipfile.ZipFile(io.BytesIO(client.get(response["download_url"]).content)) as archive:
        archive.extractall(tmp_path)
    run=subprocess.run([sys.executable,"-X","utf8","replay.py"],cwd=tmp_path/"reproduce",capture_output=True,text=True,encoding="utf-8",timeout=90)
    assert run.returncode == 0, run.stdout+run.stderr
    result=json.loads((tmp_path/"reproduce"/"replay-results.json").read_text(encoding="utf-8"))
    assert len(result["results"]) == 2 and all(r["reproduced"] for r in result["results"])
